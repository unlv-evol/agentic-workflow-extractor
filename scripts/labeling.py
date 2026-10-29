#!/usr/bin/env python3
"""Build the coders' labelling package, and import their returned workbooks.

    # one folder per coder, plus a shared calibration workbook
    python3 scripts/labeling.py build -o validation/labeling

    # a returned workbook -> the CSV sheets that score_validation.py reads
    python3 scripts/labeling.py import returned/GHAgentCI-labeling-coder-A.xlsx --coder A
    python3 scripts/labeling.py import returned/GHAgentCI-calibration.xlsx --calibration

Each coder's folder holds only that coder's workbook and the coding guide: no
detector output, nothing of the other coder's (validation/PLAN.md, Procedure).
The import refuses to write anything if a unit is missing or unknown, or an
answer is outside the codebook's vocabulary, so a typo cannot reach the scores.

Needs openpyxl (requirements-dev.txt); building the guide needs pdflatex.
"""

from __future__ import annotations

import argparse
import csv
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
VALIDATION = ROOT / "validation"
GUIDE_SRC = VALIDATION / "labeling_src" / "guide.tex"
GUIDE_PDF = "GHAgentCI-coding-guide.pdf"
DEADLINE = "3 November 2026"
CONTACT = "Daniel Ogenrwot (ogenrwot@unlv.nevada.edu)"

Q1_VALUES = ("yes", "no", "maybe")
Q2_VALUES = (
    "claude", "codex", "copilot", "gemini", "cursor", "aider", "devin", "pr_agent",
    "coderabbit", "goose", "opencode", "openhands",
    "generic_ai", "unknown", "other", "unclear",
)

# Workbook columns: (header, CSV field or None for display-only, width).
STEP_COLUMNS = [
    ("unit_id", "unit_id", 9),
    ("open workflow", None, 58),
    ("scope", "scope", 7),
    ("job_id", "job_id", 18),
    ("step_index", "step_index", 10),
    ("step_name", "step_name", 30),
    ("Q1 invocation", "q1_invocation", 14),
    ("Q2 agent", "q2_agent", 14),
    ("notes", "notes", 40),
]
FILE_COLUMNS = [
    ("unit_id", "unit_id", 9),
    ("open workflow", None, 70),
    ("Q1 invocation", "q1_invocation", 14),
    ("Q2 agent", "q2_agent", 14),
    ("notes", "notes", 40),
]
ANSWER_FIELDS = ("q1_invocation", "q2_agent", "notes")

FONT = "Arial"
HEADER_FILL = "1F3864"
INPUT_FILL = "FFF2CC"      # yellow: yours to fill
DONE_FILL = "E2EFDA"       # green: Q1 answered


def read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def link_text(permalink: str) -> str:
    """`owner/repo: path` from a GitHub blob permalink."""
    parts = permalink.split("/")
    try:
        blob = parts.index("blob")
        return f"{parts[blob - 2]}/{parts[blob - 1]}: {'/'.join(parts[blob + 2:])}"
    except (ValueError, IndexError):
        return permalink


# ---------------------------------------------------------------------------
# build


def _styles():
    from openpyxl.styles import Alignment, Font, PatternFill

    return {
        "base": Font(name=FONT, size=10),
        "bold": Font(name=FONT, size=10, bold=True),
        "title": Font(name=FONT, size=14, bold=True),
        "h2": Font(name=FONT, size=11, bold=True),
        "header": Font(name=FONT, size=10, bold=True, color="FFFFFF"),
        "link": Font(name=FONT, size=10, color="0563C1", underline="single"),
        "header_fill": PatternFill("solid", fgColor=HEADER_FILL),
        "input_fill": PatternFill("solid", fgColor=INPUT_FILL),
        "done_fill": PatternFill("solid", fgColor=DONE_FILL),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "top": Alignment(vertical="top"),
    }


def _units_sheet(wb, title: str, columns, rows: list[dict]) -> None:
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    st = _styles()
    ws = wb.create_sheet(title)
    for col, (header, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font, cell.fill = st["header"], st["header_fill"]
        ws.column_dimensions[get_column_letter(col)].width = width
    headers = [h for h, _, _ in columns]
    q1_col = get_column_letter(headers.index("Q1 invocation") + 1)
    q2_col = get_column_letter(headers.index("Q2 agent") + 1)
    notes_col = get_column_letter(headers.index("notes") + 1)

    for r, row in enumerate(rows, start=2):
        for col, (header, field, _) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=col)
            cell.font, cell.alignment = st["base"], st["top"]
            if header == "open workflow":
                cell.value = link_text(row["permalink"])
                cell.hyperlink = row["permalink"]
                cell.font = st["link"]
            elif field in ANSWER_FIELDS:
                cell.value = row.get(field) or None
                cell.fill = st["input_fill"]
                if field == "notes":
                    cell.alignment = st["wrap"]
            else:
                value = row.get(field, "")
                cell.value = int(value) if field == "step_index" and value.lstrip("-").isdigit() else (value or None)

    last = len(rows) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"

    q1 = DataValidation(type="list", formula1='"' + ",".join(Q1_VALUES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Q1",
                        error="Choose yes, no or maybe (see the Codebook tab).")
    q2 = DataValidation(type="list", formula1='"' + ",".join(Q2_VALUES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Q2",
                        error="Choose a value from the list (see the Codebook tab).")
    q1.add(f"{q1_col}2:{q1_col}{last}")
    q2.add(f"{q2_col}2:{q2_col}{last}")
    ws.add_data_validation(q1)
    ws.add_data_validation(q2)
    # Answered rows turn green; the static yellow fill stays on unanswered ones.
    ws.conditional_formatting.add(
        f"{q1_col}2:{notes_col}{last}",
        FormulaRule(formula=[f'${q1_col}2<>""'], fill=st["done_fill"]),
    )


def _write_block(ws, row: int, lines, st) -> int:
    for line in lines:
        if line is None:
            row += 1
            continue
        style, text = line
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = st[style]
        cell.alignment = st["wrap"]
        row += 1
    return row


def _table(ws, row: int, header: list[str], body: list[list[str]], st, fills=None) -> int:
    for col, text in enumerate(header, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font, cell.fill = st["header"], st["header_fill"]
    for r, values in enumerate(body, start=row + 1):
        for col, text in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=text)
            cell.font, cell.alignment = st["base"], st["wrap"]
            if fills and col in fills:
                cell.fill = st["input_fill"]
    return row + len(body) + 2


def _start_sheet(wb, who: str, tabs: list[str], joint: bool) -> None:
    st = _styles()
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 16
    ws.sheet_view.showGridLines = False
    steps = (
        [("base", "1. Code all 30 units of the Calibration tab together, discussing each disagreement."),
         ("base", "2. Settle the open rule in Section 6 of the coding guide and write the decision down."),
         ("base", "3. These units are not counted. Keep this file: it records how the rules were applied.")]
        if joint else
        [("base", "1. Read GHAgentCI-coding-guide.pdf (10 minutes). Keep it open while you code."),
         ("base", "2. Take part in the joint calibration session before coding anything here."),
         ("base", f"3. Code every row of the {' and '.join(tabs)} tab(s), on your own."),
         ("base", f"4. Save and return this file to {CONTACT} by {DEADLINE}.")]
    )
    row = _write_block(ws, 1, [
        ("title", "GHAgentCI validation: " + who),
        ("base", "For each unit, open the workflow on GitHub and answer Q1 and Q2. The coding guide (PDF) has all rules."),
        None,
        ("h2", "Steps"), *steps,
        None,
        ("h2", "Rules"),
        ("base", "Work independently: no discussion of units, and no looking at the other coder's file, until both are returned."),
        ("base", "Code it yourself: no AI assistants or scripts. Do not look up what the detector decided."),
        ("base", "Do not add, delete, rename or reorder columns. Sorting and filtering rows is fine."),
        ("base", "When in doubt, choose the best answer and write a short note."),
        None,
        ("h2", "Legend"),
    ], st)
    legend = [["yellow cell", "yours to fill: Q1, Q2 and notes"],
              ["green row", "Q1 answered"],
              ["blue link", "opens the workflow on GitHub at the commit being judged"]]
    for r, (label, meaning) in enumerate(legend, start=row):
        a = ws.cell(row=r, column=1, value=label)
        a.font = st["link"] if label.startswith("blue") else st["base"]
        if label.startswith("yellow"):
            a.fill = st["input_fill"]
        elif label.startswith("green"):
            a.fill = st["done_fill"]
        ws.cell(row=r, column=2, value=meaning).font = st["base"]
    row += len(legend) + 1
    row = _write_block(ws, row, [("h2", "Example of a completed row (not part of the sample)")], st)
    _table(ws, row,
           ["unit_id", "open workflow", "scope", "job_id", "step_index", "step_name", "Q1 invocation", "Q2 agent", "notes"],
           [["EX01", "owner/repo: .github/workflows/review.yml", "step", "review", "1", "Claude review",
             "yes", "claude", "action runs on every PR"]], st, fills={7, 8, 9})


def _codebook_sheet(wb) -> None:
    st = _styles()
    ws = wb.create_sheet("Codebook")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 95
    row = _write_block(ws, 1, [
        ("title", "Codebook (quick reference; the PDF guide is authoritative)"),
        None,
        ("h2", "Q1. Does this step (or, for scope = file and the Recall tab, this workflow) invoke a coding agent when the workflow runs?"),
        ("base", "A coding agent is an LLM-based tool that acts on the repository or its activity: it reads code, diffs, issues or comments, "
                 "and produces reviews, comments, labels, commits or pull requests."),
    ], st)
    row = _table(ws, row + 1, ["answer", "use when"], [
        ["yes", "It runs such an agent: a vendor agent action, an agent CLI (claude -p, codex exec), a wrapper action that runs an agent, "
                "copilot-setup-steps.yml, or a gh-aw workflow."],
        ["maybe", "An LLM is called but it is unclear it acts as an agent (e.g. only summarises text), or the agent step cannot run "
                  "(if: false, dry-run). Also: link no longer resolves (note 'link dead')."],
        ["no", "No agent runs: only installs an agent, a key feeding tests, a name in echo/comment/docs, an unrelated tool with the "
               "same name, or not a working workflow."],
    ], st)
    row = _write_block(ws, row, [
        ("h2", "Q2. Which agent? Only when Q1 is yes or maybe; leave empty for no. Spelling must match the list."),
    ], st)
    _table(ws, row + 1, ["value", "use when"], [
        ["claude, codex, copilot, gemini", "the vendor's agent, directly or through a wrapper that clearly uses it"],
        ["cursor, aider, devin, pr_agent", "pr_agent is Qodo's PR-Agent (Codium-ai/pr-agent)"],
        ["coderabbit, goose, opencode, openhands", ""],
        ["generic_ai", "a third-party AI action not tied to one of the vendors above"],
        ["unknown", "a gh-aw file that does not say which engine it uses"],
        ["other", "a coding agent not in this list"],
        ["unclear", "cannot be determined from the workflow"],
    ], st)


def _workbook(path: Path, who: str, sheets: list[tuple[str, list, list[dict]]], joint: bool) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    _start_sheet(wb, who, [title for title, _, _ in sheets], joint)
    _codebook_sheet(wb)
    for title, columns, rows in sheets:
        _units_sheet(wb, title, columns, rows)
    wb.save(path)


def build_guide(out: Path) -> Path:
    """Compile the guide with pdflatex, twice for stable layout."""
    tex = shutil.which("pdflatex") or "/Library/TeX/texbin/pdflatex"
    work = GUIDE_SRC.parent
    for _ in range(2):
        subprocess.run([tex, "-interaction=nonstopmode", "-halt-on-error", GUIDE_SRC.name],
                       cwd=work, check=True, capture_output=True)
    pdf = out / GUIDE_PDF
    shutil.copyfile(work / "guide.pdf", pdf)
    return pdf


def build(out: Path, source: Path) -> None:
    precision = {c: read_csv(source / f"precision_sheet_{c}.csv") for c in "AB"}
    recall = {c: read_csv(source / f"recall_sheet_{c}.csv") for c in "AB"}
    calibration = read_csv(source / "calibration_sheet.csv")
    for kind, sheets in (("precision", precision), ("recall", recall)):
        ids = [[r["unit_id"] for r in sheets[c]] for c in "AB"]
        if ids[0] != ids[1]:
            raise SystemExit(f"{kind}: coder A and B sheets hold different units")
        answered = sum(1 for c in "AB" for r in sheets[c] if (r.get("q1_invocation") or "").strip())
        if answered:
            raise SystemExit(f"{kind}: sheets already hold {answered} answers; refusing to overwrite")

    out.mkdir(parents=True, exist_ok=True)
    guide = build_guide(out)
    for coder in "AB":
        folder = out / f"coder-{coder}"
        folder.mkdir(exist_ok=True)
        _workbook(folder / f"GHAgentCI-labeling-coder-{coder}.xlsx", f"coder {coder}",
                  [("Precision", STEP_COLUMNS, precision[coder]),
                   ("Recall", FILE_COLUMNS, recall[coder])], joint=False)
        shutil.copyfile(guide, folder / GUIDE_PDF)
    shared = out / "shared"
    shared.mkdir(exist_ok=True)
    _workbook(shared / "GHAgentCI-calibration.xlsx", "calibration (both coders, together)",
              [("Calibration", STEP_COLUMNS, calibration)], joint=True)
    shutil.copyfile(guide, shared / GUIDE_PDF)
    guide.unlink()
    print(f"built {out}: coder-A, coder-B, shared", file=sys.stderr)


# ---------------------------------------------------------------------------
# import


def read_answers(path: Path, tab: str) -> dict[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if tab not in wb.sheetnames:
        raise SystemExit(f"{path.name}: no '{tab}' tab")
    rows = wb[tab].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    index = {name: header.index(name) for name in ("unit_id", "Q1 invocation", "Q2 agent", "notes")
             if name in header}
    if len(index) < 4:
        raise SystemExit(f"{path.name} [{tab}]: columns were renamed or removed; expected unit_id, "
                         "Q1 invocation, Q2 agent, notes")
    answers = {}
    for values in rows:
        unit = values[index["unit_id"]]
        if unit is None:
            continue

        def text(col):
            value = values[index[col]]
            return "" if value is None else str(value).strip()

        answers[str(unit).strip()] = {
            "q1_invocation": text("Q1 invocation").lower(),
            "q2_agent": text("Q2 agent").lower(),
            "notes": text("notes"),
        }
    return answers


def merge(csv_path: Path, answers: dict[str, dict], label: str) -> tuple[list[dict], list[str], list[str]]:
    """The CSV rows with answers filled in, plus errors and warnings."""
    rows = read_csv(csv_path)
    known = {r["unit_id"] for r in rows}
    errors, warnings = [], []
    missing = sorted(known - set(answers))
    extra = sorted(set(answers) - known)
    if missing:
        errors.append(f"{label}: {len(missing)} unit(s) missing from the workbook, e.g. {missing[:5]}")
    if extra:
        errors.append(f"{label}: {len(extra)} unknown unit(s), e.g. {extra[:5]}")
    for unit, a in answers.items():
        if a["q1_invocation"] and a["q1_invocation"] not in Q1_VALUES:
            errors.append(f"{label} {unit}: Q1 '{a['q1_invocation']}' is not yes/no/maybe")
        if a["q2_agent"] and a["q2_agent"] not in Q2_VALUES:
            errors.append(f"{label} {unit}: Q2 '{a['q2_agent']}' is not in the codebook")
        if a["q1_invocation"] == "no" and a["q2_agent"]:
            warnings.append(f"{label} {unit}: Q2 given although Q1 is no (ignored by the scorer)")
        if a["q1_invocation"] in ("yes", "maybe") and not a["q2_agent"]:
            warnings.append(f"{label} {unit}: Q1 is {a['q1_invocation']} but Q2 is empty")
    for row in rows:
        row.update(answers.get(row["unit_id"], {}))
    return rows, errors, warnings


def import_workbook(path: Path, coder: str | None, calibration: bool, target: Path) -> int:
    jobs = ([("Calibration", target / "calibration_sheet.csv")] if calibration else
            [("Precision", target / f"precision_sheet_{coder}.csv"),
             ("Recall", target / f"recall_sheet_{coder}.csv")])
    merged, errors, warnings = [], [], []
    for tab, csv_path in jobs:
        rows, e, w = merge(csv_path, read_answers(path, tab), tab)
        merged.append((csv_path, rows))
        errors += e
        warnings += w
    for message in warnings[:30]:
        print(f"  warning: {message}", file=sys.stderr)
    if len(warnings) > 30:
        print(f"  ... and {len(warnings) - 30} more warnings", file=sys.stderr)
    if errors:
        for message in errors[:30]:
            print(f"  ERROR: {message}", file=sys.stderr)
        print("nothing written; fix the workbook and import again", file=sys.stderr)
        return 2
    for csv_path, rows in merged:
        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        done = sum(1 for r in rows if r["q1_invocation"])
        print(f"wrote {csv_path.name}: {done}/{len(rows)} units answered", file=sys.stderr)
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = parser.add_subparsers(dest="mode", required=True)
    b = sub.add_parser("build", help="write the package: coder-A/, coder-B/, shared/")
    b.add_argument("-o", "--out-dir", default=str(VALIDATION / "labeling"))
    b.add_argument("--source", default=str(VALIDATION), help="where the blank CSV sheets are")
    i = sub.add_parser("import", help="read a returned workbook into the CSV sheets")
    i.add_argument("workbook")
    who = i.add_mutually_exclusive_group(required=True)
    who.add_argument("--coder", choices=["A", "B"])
    who.add_argument("--calibration", action="store_true")
    i.add_argument("--target", default=str(VALIDATION), help="where the CSV sheets are")
    args = parser.parse_args(argv)

    if args.mode == "build":
        build(Path(args.out_dir), Path(args.source))
        return 0
    return import_workbook(Path(args.workbook), args.coder, args.calibration, Path(args.target))


if __name__ == "__main__":
    raise SystemExit(main())
